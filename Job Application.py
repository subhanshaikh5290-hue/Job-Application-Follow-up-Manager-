import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
from datetime import datetime
import os

FILE_NAME = "job_applications.xlsx"

# ---------------- Excel Setup ----------------
def setup_excel():
    if not os.path.exists(FILE_NAME):
        wb = Workbook()
        ws = wb.active
        ws.title = "Applications"

        ws.append([
            "ID",
            "Company",
            "Position",
            "Application Date",
            "Follow-up Date",
            "Status",
            "Notes"
        ])

        wb.save(FILE_NAME)


# ---------------- Clear Fields ----------------
def clear_fields():
    company_entry.delete(0, tk.END)
    position_entry.delete(0, tk.END)
    application_entry.delete(0, tk.END)
    followup_entry.delete(0, tk.END)
    status_combo.set("Applied")
    notes_entry.delete(0, tk.END)


# ---------------- Add Application ----------------
def add_application():
    company = company_entry.get().strip()
    position = position_entry.get().strip()
    application_date = application_entry.get().strip()
    followup_date = followup_entry.get().strip()
    status = status_combo.get()
    notes = notes_entry.get().strip()

    if not company or not position or not application_date:
        messagebox.showwarning(
            "Missing Data",
            "Please enter Company, Position and Application Date."
        )
        return

    try:
        datetime.strptime(application_date, "%d-%m-%Y")

        if followup_date:
            datetime.strptime(followup_date, "%d-%m-%Y")

    except ValueError:
        messagebox.showerror(
            "Invalid Date",
            "Date format must be DD-MM-YYYY."
        )
        return

    wb = load_workbook(FILE_NAME)
    ws = wb["Applications"]

    new_id = ws.max_row

    ws.append([
        new_id,
        company,
        position,
        application_date,
        followup_date,
        status,
        notes
    ])

    wb.save(FILE_NAME)

    messagebox.showinfo(
        "Success",
        "Job application added successfully!"
    )

    clear_fields()
    show_all()


# ---------------- Show All ----------------
def show_all():
    for item in tree.get_children():
        tree.delete(item)

    wb = load_workbook(FILE_NAME)
    ws = wb["Applications"]

    for row in ws.iter_rows(min_row=2, values_only=True):
        tree.insert("", tk.END, values=row)


# ---------------- Search ----------------
def search_application():
    search_text = search_entry.get().strip().lower()

    for item in tree.get_children():
        tree.delete(item)

    wb = load_workbook(FILE_NAME)
    ws = wb["Applications"]

    for row in ws.iter_rows(min_row=2, values_only=True):
        company = str(row[1]).lower()
        position = str(row[2]).lower()

        if search_text in company or search_text in position:
            tree.insert("", tk.END, values=row)


# ---------------- Delete ----------------
def delete_application():
    selected = tree.selection()

    if not selected:
        messagebox.showwarning(
            "Select Application",
            "Please select an application first."
        )
        return

    item = tree.item(selected[0])
    application_id = item["values"][0]

    wb = load_workbook(FILE_NAME)
    ws = wb["Applications"]

    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value == application_id:
            ws.delete_rows(row, 1)
            break

    wb.save(FILE_NAME)

    messagebox.showinfo(
        "Deleted",
        "Application deleted successfully!"
    )

    show_all()


# ---------------- Load Selected Data ----------------
def load_selected():
    selected = tree.selection()

    if not selected:
        messagebox.showwarning(
            "Select Application",
            "Please select an application."
        )
        return

    values = tree.item(selected[0])["values"]

    clear_fields()

    company_entry.insert(0, values[1])
    position_entry.insert(0, values[2])
    application_entry.insert(0, values[3])
    followup_entry.insert(0, values[4])
    status_combo.set(values[5])
    notes_entry.insert(0, values[6])


# ---------------- Update ----------------
def update_application():
    selected = tree.selection()

    if not selected:
        messagebox.showwarning(
            "Select Application",
            "Please select an application."
        )
        return

    values = tree.item(selected[0])["values"]
    application_id = values[0]

    company = company_entry.get().strip()
    position = position_entry.get().strip()
    application_date = application_entry.get().strip()
    followup_date = followup_entry.get().strip()
    status = status_combo.get()
    notes = notes_entry.get().strip()

    if not company or not position or not application_date:
        messagebox.showwarning(
            "Missing Data",
            "Please enter required details."
        )
        return

    wb = load_workbook(FILE_NAME)
    ws = wb["Applications"]

    for row in range(2, ws.max_row + 1):

        if ws.cell(row, 1).value == application_id:

            ws.cell(row, 2).value = company
            ws.cell(row, 3).value = position
            ws.cell(row, 4).value = application_date
            ws.cell(row, 5).value = followup_date
            ws.cell(row, 6).value = status
            ws.cell(row, 7).value = notes

            break

    wb.save(FILE_NAME)

    messagebox.showinfo(
        "Updated",
        "Application updated successfully!"
    )

    clear_fields()
    show_all()


# ---------------- GUI ----------------
setup_excel()

root = tk.Tk()
root.title("Job Application Follow-up Manager")
root.geometry("1050x650")
root.resizable(False, False)

title = tk.Label(
    root,
    text="JOB APPLICATION FOLLOW-UP MANAGER",
    font=("Arial", 22, "bold")
)

title.pack(pady=15)


# -------- Input Frame --------
input_frame = tk.Frame(root)
input_frame.pack(pady=5)

tk.Label(input_frame, text="Company Name").grid(
    row=0, column=0, padx=10, pady=5
)

company_entry = tk.Entry(input_frame, width=25)
company_entry.grid(row=0, column=1, padx=10)


tk.Label(input_frame, text="Job Position").grid(
    row=0, column=2, padx=10
)

position_entry = tk.Entry(input_frame, width=25)
position_entry.grid(row=0, column=3, padx=10)


tk.Label(input_frame, text="Application Date").grid(
    row=1, column=0, padx=10, pady=5
)

application_entry = tk.Entry(input_frame, width=25)
application_entry.grid(row=1, column=1, padx=10)

application_entry.insert(
    0,
    datetime.now().strftime("%d-%m-%Y")
)


tk.Label(input_frame, text="Follow-up Date").grid(
    row=1, column=2, padx=10
)

followup_entry = tk.Entry(input_frame, width=25)
followup_entry.grid(row=1, column=3, padx=10)


tk.Label(input_frame, text="Status").grid(
    row=2, column=0, padx=10, pady=5
)

status_combo = ttk.Combobox(
    input_frame,
    values=[
        "Applied",
        "Interview",
        "Selected",
        "Rejected"
    ],
    width=22,
    state="readonly"
)

status_combo.grid(row=2, column=1)
status_combo.set("Applied")


tk.Label(input_frame, text="Notes").grid(
    row=2, column=2, padx=10
)

notes_entry = tk.Entry(input_frame, width=25)
notes_entry.grid(row=2, column=3)


# -------- Buttons --------
button_frame = tk.Frame(root)
button_frame.pack(pady=15)

tk.Button(
    button_frame,
    text="Add Application",
    width=18,
    command=add_application
).grid(row=0, column=0, padx=5)

tk.Button(
    button_frame,
    text="Load Selected",
    width=18,
    command=load_selected
).grid(row=0, column=1, padx=5)

tk.Button(
    button_frame,
    text="Update",
    width=18,
    command=update_application
).grid(row=0, column=2, padx=5)

tk.Button(
    button_frame,
    text="Delete",
    width=18,
    command=delete_application
).grid(row=0, column=3, padx=5)

tk.Button(
    button_frame,
    text="Clear",
    width=18,
    command=clear_fields
).grid(row=0, column=4, padx=5)


# -------- Search --------
search_frame = tk.Frame(root)
search_frame.pack(pady=5)

tk.Label(
    search_frame,
    text="Search Company / Position:"
).pack(side=tk.LEFT)

search_entry = tk.Entry(
    search_frame,
    width=30
)

search_entry.pack(side=tk.LEFT, padx=10)

tk.Button(
    search_frame,
    text="Search",
    command=search_application
).pack(side=tk.LEFT)

tk.Button(
    search_frame,
    text="Show All",
    command=show_all
).pack(side=tk.LEFT, padx=5)


# -------- Table --------
table_frame = tk.Frame(root)
table_frame.pack(pady=15)

columns = (
    "ID",
    "Company",
    "Position",
    "Application Date",
    "Follow-up Date",
    "Status",
    "Notes"
)

tree = ttk.Treeview(
    table_frame,
    columns=columns,
    show="headings",
    height=12
)

for col in columns:
    tree.heading(col, text=col)
    tree.column(col, width=130)

tree.pack(side=tk.LEFT)

scrollbar = ttk.Scrollbar(
    table_frame,
    orient="vertical",
    command=tree.yview
)

tree.configure(yscrollcommand=scrollbar.set)
scrollbar.pack(side=tk.RIGHT, fill=tk.Y)


show_all()

root.mainloop()